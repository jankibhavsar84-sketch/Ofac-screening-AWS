from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from .models import EntityExample


@dataclass(frozen=True)
class BatchRowMeta:
    key: str
    display_name: str
    ui_type: str


@dataclass(frozen=True)
class ParsedBatchUpload:
    queries: dict[str, EntityExample]
    row_meta: list[BatchRowMeta]


class BatchUploadValidationError(ValueError):
    def __init__(self, errors: list[str]) -> None:
        self.errors = errors
        preview = " ".join(errors[:4])
        remaining = f" (+{len(errors) - 4} more)" if len(errors) > 4 else ""
        super().__init__(f"Upload validation failed. {preview}{remaining}")


_COUNTRY_CODE_MAP = {
    "united states": "us",
    "usa": "us",
    "america": "us",
    "india": "in",
    "united kingdom": "gb",
    "uk": "gb",
    "canada": "ca",
}


def _safe_trim(value: Any) -> str:
    return str(value or "").strip()


def _normalize_header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", _safe_trim(value).lower())


def _normalize_customer_type(value: Any) -> str:
    lowered = _safe_trim(value).lower()
    if lowered in {"entity", "organization", "company", "e"}:
        return "Entity"
    return "Person"


def _get_by_aliases(values: dict[str, Any], aliases: list[str]) -> str:
    for alias in aliases:
        found = _safe_trim(values.get(_normalize_header_key(alias)))
        if found:
            return found
    return ""


def _map_parsed_row(values: dict[str, Any]) -> dict[str, str]:
    customer_type_raw = _get_by_aliases(values, ["customerType"])
    return {
        "partyKey": _get_by_aliases(values, ["partyKey", "party_key", "party key"]),
        "customerType": _normalize_customer_type(
            customer_type_raw or _get_by_aliases(values, ["partyType", "party type", "party_type"])
        ),
        "customerTypeRaw": customer_type_raw,
        "partyType": _get_by_aliases(
            values,
            ["partyType", "party type", "party_type", "partyTypeCode", "party type code"],
        ),
        "title": _get_by_aliases(values, ["title"]),
        "gender": _get_by_aliases(values, ["gender", "genderCode", "gender code", "gender_code"]),
        "primaryFirstName": _get_by_aliases(values, ["primaryFirstName"]),
        "primaryMiddleName": _get_by_aliases(values, ["primaryMiddleName"]),
        "primaryLastName": _get_by_aliases(values, ["primaryLastName"]),
        "primaryMaidenName": _get_by_aliases(values, ["primaryMaidenName"]),
        "primaryFullName": _get_by_aliases(values, ["primaryFullName"]),
        "alias1FirstName": _get_by_aliases(values, ["alias1FirstName"]),
        "alias1MiddleName": _get_by_aliases(values, ["alias1MiddleName"]),
        "alias1LastName": _get_by_aliases(values, ["alias1LastName"]),
        "alias1MaidenName": _get_by_aliases(values, ["alias1MaidenName"]),
        "alias1FullName": _get_by_aliases(values, ["alias1FullName"]),
        "alias2FirstName": _get_by_aliases(values, ["alias2FirstName"]),
        "alias2MiddleName": _get_by_aliases(values, ["alias2MiddleName"]),
        "alias2LastName": _get_by_aliases(values, ["alias2LastName"]),
        "alias2MaidenName": _get_by_aliases(values, ["alias2MaidenName"]),
        "alias2FullName": _get_by_aliases(values, ["alias2FullName"]),
        "alias3FirstName": _get_by_aliases(values, ["alias3FirstName"]),
        "alias3MiddleName": _get_by_aliases(values, ["alias3MiddleName"]),
        "alias3LastName": _get_by_aliases(values, ["alias3LastName"]),
        "alias3MaidenName": _get_by_aliases(values, ["alias3MaidenName"]),
        "alias3FullName": _get_by_aliases(values, ["alias3FullName"]),
        "partyId1Type": _get_by_aliases(values, ["partyId1Type"]),
        "partyId1Value": _get_by_aliases(values, ["partyId1Value"]),
        "partyId1IDCountry": _get_by_aliases(values, ["partyId1IDCountry"]),
        "partyId2Type": _get_by_aliases(values, ["partyId2Type"]),
        "partyId2Value": _get_by_aliases(values, ["partyId2Value"]),
        "partyId2IDCountry": _get_by_aliases(values, ["partyId2IDCountry"]),
        "partyId3Type": _get_by_aliases(values, ["partyId3Type"]),
        "partyId3Value": _get_by_aliases(values, ["partyId3Value"]),
        "partyId3IDCountry": _get_by_aliases(values, ["partyId3IDCountry"]),
        "address1Line1": _get_by_aliases(values, ["address1Line1"]),
        "address1Line2": _get_by_aliases(values, ["address1Line2"]),
        "address1City": _get_by_aliases(values, ["address1City"]),
        "address1ZipCode": _get_by_aliases(values, ["address1ZipCode"]),
        "address1Country": _get_by_aliases(values, ["address1Country"]),
        "address1stateProvince": _get_by_aliases(values, ["address1stateProvince"]),
        "address2Line1": _get_by_aliases(values, ["address2Line1"]),
        "address2Line2": _get_by_aliases(values, ["address2Line2"]),
        "address2City": _get_by_aliases(values, ["address2City"]),
        "address2ZipCode": _get_by_aliases(values, ["address2ZipCode"]),
        "address2Country": _get_by_aliases(values, ["address2Country"]),
        "address2stateProvince": _get_by_aliases(values, ["address2stateProvince"]),
        "address3Line1": _get_by_aliases(values, ["address3Line1"]),
        "address3Line2": _get_by_aliases(values, ["address3Line2"]),
        "address3City": _get_by_aliases(values, ["address3City"]),
        "address3ZipCode": _get_by_aliases(values, ["address3ZipCode"]),
        "address3Country": _get_by_aliases(values, ["address3Country"]),
        "address3stateProvince": _get_by_aliases(values, ["address3stateProvince"]),
        "birthCountry": _get_by_aliases(values, ["birthCountry"]),
        "birthLocation": _get_by_aliases(values, ["birthLocation"]),
        "nationalityCountry1": _get_by_aliases(values, ["nationalityCountry1"]),
        "nationalityCountry2": _get_by_aliases(values, ["nationalityCountry2"]),
        "nationalityCountry3": _get_by_aliases(values, ["nationalityCountry3"]),
        "yearOfBirth": _get_by_aliases(values, ["yearOfBirth"]),
        "firstName": _get_by_aliases(values, ["firstName", "primaryFirstName", "first name"]),
        "middleName": _get_by_aliases(values, ["middleName", "primaryMiddleName", "middle name"]),
        "lastName": _get_by_aliases(values, ["lastName", "primaryLastName", "last name"]),
        "fullName": _get_by_aliases(values, ["fullName", "primaryFullName", "primary name", "name"]),
        "aliasName": _get_by_aliases(values, ["aliasName", "alias1FullName", "alias", "aka"]),
        "addressLine1": _get_by_aliases(values, ["addressLine1", "address1Line1", "address line 1"]),
        "addressLine2": _get_by_aliases(values, ["addressLine2", "address1Line2", "address line 2"]),
        "city": _get_by_aliases(values, ["city", "address1City"]),
        "state": _get_by_aliases(values, ["state", "address1stateProvince", "stateProvince"]),
        "zip": _get_by_aliases(values, ["zip", "zipCode", "address1ZipCode"]),
        "country": _get_by_aliases(values, ["country", "address1Country", "nationalityCountry1", "birthCountry"]),
        "addresses": _get_by_aliases(values, ["addresses"]),
        "countries": _get_by_aliases(values, ["countries"]),
        "idType": _get_by_aliases(values, ["idType", "partyId1Type", "id code", "idCode"]),
        "idCode": _get_by_aliases(values, ["idCode", "partyId1Type", "id type"]),
        "idNumber": _get_by_aliases(values, ["idNumber", "partyId1Value", "id value"]),
        "idCountry": _get_by_aliases(values, ["idCountry", "partyId1IDCountry", "id issue country"]),
        "idIssueCountry": _get_by_aliases(values, ["idIssueCountry", "partyId1IDCountry", "id country"]),
        "countryOfBirth": _get_by_aliases(values, ["countryOfBirth", "birthCountry"]),
        "dateOfBirth": _get_by_aliases(values, ["dateOfBirth", "DateOfBirth"]),
        "countryOfCitizenship": _get_by_aliases(values, ["countryOfCitizenship", "nationalityCountry1"]),
    }


def _parse_csv_rows(body: bytes) -> list[dict[str, str]]:
    text = body.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if len(rows) < 2:
        return []

    headers = [_normalize_header_key(cell) for cell in rows[0]]
    parsed: list[dict[str, str]] = []
    for raw_row in rows[1:]:
        if not any(_safe_trim(cell) for cell in raw_row):
            continue
        values: dict[str, str] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            values[header] = _safe_trim(raw_row[index] if index < len(raw_row) else "")
        parsed.append(_map_parsed_row(values))
    return parsed


def _parse_xlsx_rows(body: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    worksheet = workbook.worksheets[0] if workbook.worksheets else None
    if worksheet is None:
        return []

    iterator = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration:
        return []

    headers = [_normalize_header_key(cell) for cell in header_row]
    if not any(headers):
        return []

    parsed: list[dict[str, str]] = []
    for raw_row in iterator:
        if not any(_safe_trim(cell) for cell in raw_row):
            continue
        values: dict[str, str] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            values[header] = _safe_trim(raw_row[index] if index < len(raw_row) else "")
        parsed.append(_map_parsed_row(values))
    return parsed


def _parse_comma_values(value: Any) -> list[str]:
    raw = _safe_trim(value)
    if not raw:
        return []
    return [_safe_trim(token) for token in raw.split(",") if _safe_trim(token)]


def _build_structured_name(first: str, middle: str, last: str, maiden: str, full_name: str) -> str:
    split_name = " ".join(part for part in [_safe_trim(first), _safe_trim(middle), _safe_trim(last), _safe_trim(maiden)] if part)
    return _safe_trim(full_name) or split_name


def _build_address_line(
    line1: Any,
    line2: Any,
    city: Any,
    state_province: Any,
    zip_code: Any,
    country: Any,
) -> str:
    return ", ".join(
        part
        for part in [
            _safe_trim(line1),
            _safe_trim(line2),
            _safe_trim(city),
            _safe_trim(state_province),
            _safe_trim(zip_code),
            _safe_trim(country),
        ]
        if part
    )


def _to_country_code(value: str) -> str:
    normalized = _safe_trim(value).lower()
    if not normalized:
        return ""
    if len(normalized) == 2:
        return normalized
    return _COUNTRY_CODE_MAP.get(normalized, normalized)


def _ui_type_to_schema(ui_type: str) -> str:
    if ui_type == "Individual":
        return "Person"
    if ui_type == "Organization":
        return "Company"
    if ui_type == "Unknown":
        return "Unknown"
    if ui_type == "Vessel":
        return "Vessel"
    return "Aircraft"


def _build_entity_example(item: dict[str, Any]) -> EntityExample:
    schema = _ui_type_to_schema(str(item.get("uiType") or "Unknown"))
    name_values: list[str] = []
    if item.get("uiType") == "Individual":
        split_name = " ".join(
            part
            for part in [
                _safe_trim(item.get("firstName")),
                _safe_trim(item.get("middleName")),
                _safe_trim(item.get("lastName")),
            ]
            if part
        )
        full_name = _safe_trim(item.get("fullName"))
        if split_name:
            name_values.append(split_name)
        if full_name and full_name.lower() != split_name.lower():
            name_values.append(full_name)
    else:
        full_name = _safe_trim(item.get("fullName"))
        if full_name:
            name_values.append(full_name)

    properties: dict[str, Any] = {"name": name_values}
    alias_name = _safe_trim(item.get("aliasName"))
    if alias_name:
        properties["alias"] = [alias_name]

    date_of_birth = _safe_trim(item.get("dateOfBirth"))
    if item.get("uiType") == "Individual" and date_of_birth:
        properties["birthDate"] = [date_of_birth]

    countries = [_to_country_code(value) for value in item.get("countries", []) if _to_country_code(value)]
    if countries:
        if schema == "Person":
            properties["nationality"] = countries
        else:
            properties["country"] = countries

    addresses = [_safe_trim(value) for value in item.get("addresses", []) if _safe_trim(value)]
    if addresses:
        properties["address"] = addresses

    normalized_ids = [
        {
            "idType": _safe_trim(doc.get("idType")),
            "idValue": _safe_trim(doc.get("idNumber")),
            "idCountry": _safe_trim(doc.get("idCountry")),
        }
        for doc in item.get("ids", [])
        if _safe_trim(doc.get("idNumber"))
    ]
    id_numbers = [_safe_trim(doc.get("idValue")) for doc in normalized_ids if _safe_trim(doc.get("idValue"))]
    if id_numbers:
        if schema == "Person":
            properties["idNumber"] = id_numbers
        else:
            properties["registrationNumber"] = id_numbers
        properties["ids"] = normalized_ids

    birth_location = _safe_trim(item.get("birthLocation"))
    if birth_location:
        properties["birthLocation"] = [birth_location]
    gender = _safe_trim(item.get("gender"))
    if gender:
        properties["gender"] = [gender]
    title = _safe_trim(item.get("title"))
    if title:
        properties["title"] = [title]

    return EntityExample(schema=schema, properties=properties)


def parse_batch_upload(filename: str, body: bytes) -> ParsedBatchUpload:
    safe_name = _safe_trim(filename).lower()
    if safe_name.endswith(".csv"):
        rows = _parse_csv_rows(body)
    elif safe_name.endswith(".xlsx"):
        rows = _parse_xlsx_rows(body)
    else:
        raise BatchUploadValidationError(["Only CSV or XLSX files are allowed."])

    if not rows:
        raise BatchUploadValidationError(["No rows found in the file."])

    queries: dict[str, EntityExample] = {}
    row_meta: list[BatchRowMeta] = []
    validation_errors: list[str] = []
    seen_party_keys: set[str] = set()

    for index, row in enumerate(rows):
        row_number = index + 2
        party_key = _safe_trim(row.get("partyKey"))
        normalized_party_key = party_key.upper()
        if not party_key:
            validation_errors.append(f"Row {row_number}: PartyKey is required.")
            continue
        if normalized_party_key in seen_party_keys:
            validation_errors.append(
                f"Row {row_number}: Duplicate PartyKey '{party_key}'. PartyKey must be unique within the file."
            )
            continue
        seen_party_keys.add(normalized_party_key)

        raw_party_type = _safe_trim(row.get("partyType"))
        normalized_party_type = raw_party_type.upper()
        has_legacy_customer_type = bool(_safe_trim(row.get("customerTypeRaw")))
        customer_type = "Entity" if row.get("customerType") == "Entity" else "Person"

        if normalized_party_type == "I":
            ui_type = "Individual"
        elif normalized_party_type == "E":
            ui_type = "Organization"
        elif raw_party_type:
            ui_type = "Unknown"
        elif has_legacy_customer_type:
            ui_type = "Organization" if customer_type == "Entity" else "Individual"
        elif _safe_trim(row.get("firstName")) or _safe_trim(row.get("lastName")):
            ui_type = "Individual"
        elif _safe_trim(row.get("fullName")):
            ui_type = "Organization"
        else:
            ui_type = "Unknown"

        raw_gender = _safe_trim(row.get("gender"))
        if raw_gender and raw_gender.upper() not in {"M", "F"}:
            validation_errors.append(f"Row {row_number}: Gender code must be M, F, or blank.")
            continue

        first_name = _safe_trim(row.get("primaryFirstName") or row.get("firstName"))
        middle_name = _safe_trim(row.get("primaryMiddleName") or row.get("middleName"))
        last_name = _safe_trim(row.get("primaryLastName") or row.get("lastName"))
        maiden_name = _safe_trim(row.get("primaryMaidenName"))
        split_name = " ".join(part for part in [first_name, middle_name, last_name] if part)
        explicit_full_name = _build_structured_name(
            first_name,
            middle_name,
            last_name,
            maiden_name,
            _safe_trim(row.get("primaryFullName") or row.get("fullName")),
        )
        has_split_first_last = bool(first_name and last_name)
        has_full_name = bool(explicit_full_name)
        full_name = explicit_full_name or (split_name if ui_type == "Individual" else "")

        if ui_type == "Individual" and not has_full_name and not has_split_first_last:
            validation_errors.append(
                f"Row {row_number}: For Individual, provide either Full Name or both First Name and Last Name."
            )
            continue
        if ui_type != "Individual" and not full_name:
            validation_errors.append(f"Row {row_number}: For Organization or Unknown, Full Name is required.")
            continue

        addresses = _parse_comma_values(row.get("addresses"))
        address_candidates = [
            _build_address_line(
                row.get("address1Line1"),
                row.get("address1Line2"),
                row.get("address1City"),
                row.get("address1stateProvince"),
                row.get("address1ZipCode"),
                row.get("address1Country"),
            ),
            _build_address_line(
                row.get("address2Line1"),
                row.get("address2Line2"),
                row.get("address2City"),
                row.get("address2stateProvince"),
                row.get("address2ZipCode"),
                row.get("address2Country"),
            ),
            _build_address_line(
                row.get("address3Line1"),
                row.get("address3Line2"),
                row.get("address3City"),
                row.get("address3stateProvince"),
                row.get("address3ZipCode"),
                row.get("address3Country"),
            ),
            _build_address_line(
                row.get("addressLine1"),
                row.get("addressLine2"),
                row.get("city"),
                row.get("state"),
                row.get("zip"),
                row.get("country"),
            ),
        ]
        for address in [_safe_trim(value) for value in address_candidates if _safe_trim(value)]:
            if address not in addresses:
                addresses.append(address)
        if not addresses:
            one_line_address = ", ".join(
                value
                for value in [
                    _safe_trim(row.get("addressLine1")),
                    _safe_trim(row.get("addressLine2")),
                    _safe_trim(row.get("city")),
                    _safe_trim(row.get("state")),
                    _safe_trim(row.get("zip")),
                ]
                if value
            )
            if one_line_address:
                addresses.append(one_line_address)

        countries = list(
            dict.fromkeys(
                value
                for value in [
                    *_parse_comma_values(row.get("countries")),
                    _safe_trim(row.get("country")),
                    _safe_trim(row.get("countryOfCitizenship")),
                    _safe_trim(row.get("address1Country")),
                    _safe_trim(row.get("address2Country")),
                    _safe_trim(row.get("address3Country")),
                    _safe_trim(row.get("birthCountry")),
                    _safe_trim(row.get("nationalityCountry1")),
                    _safe_trim(row.get("nationalityCountry2")),
                    _safe_trim(row.get("nationalityCountry3")),
                ]
                if value
            )
        )

        alias_1 = _build_structured_name(
            _safe_trim(row.get("alias1FirstName")),
            _safe_trim(row.get("alias1MiddleName")),
            _safe_trim(row.get("alias1LastName")),
            _safe_trim(row.get("alias1MaidenName")),
            _safe_trim(row.get("alias1FullName")),
        )
        alias_2 = _build_structured_name(
            _safe_trim(row.get("alias2FirstName")),
            _safe_trim(row.get("alias2MiddleName")),
            _safe_trim(row.get("alias2LastName")),
            _safe_trim(row.get("alias2MaidenName")),
            _safe_trim(row.get("alias2FullName")),
        )
        alias_3 = _build_structured_name(
            _safe_trim(row.get("alias3FirstName")),
            _safe_trim(row.get("alias3MiddleName")),
            _safe_trim(row.get("alias3LastName")),
            _safe_trim(row.get("alias3MaidenName")),
            _safe_trim(row.get("alias3FullName")),
        )
        merged_alias = ", ".join(
            value for value in [_safe_trim(row.get("aliasName")), alias_1, alias_2, alias_3] if value
        )

        ids = [
            {
                "idType": _safe_trim(row.get("partyId1Type") or row.get("idType") or row.get("idCode")),
                "idNumber": _safe_trim(row.get("partyId1Value") or row.get("idNumber")),
                "idCountry": _safe_trim(
                    row.get("partyId1IDCountry") or row.get("idCountry") or row.get("idIssueCountry")
                ),
            },
            {
                "idType": _safe_trim(row.get("partyId2Type")),
                "idNumber": _safe_trim(row.get("partyId2Value")),
                "idCountry": _safe_trim(row.get("partyId2IDCountry")),
            },
            {
                "idType": _safe_trim(row.get("partyId3Type")),
                "idNumber": _safe_trim(row.get("partyId3Value")),
                "idCountry": _safe_trim(row.get("partyId3IDCountry")),
            },
        ]
        ids = [doc for doc in ids if _safe_trim(doc.get("idNumber"))]

        item = {
            "uiType": ui_type,
            "firstName": first_name,
            "middleName": middle_name,
            "lastName": last_name,
            "fullName": full_name,
            "aliasName": merged_alias,
            "dateOfBirth": _safe_trim(row.get("dateOfBirth") or row.get("yearOfBirth")),
            "countries": countries,
            "addresses": addresses,
            "ids": ids,
            "birthLocation": _safe_trim(row.get("birthLocation") or row.get("birthCountry")),
            "gender": raw_gender,
            "title": _safe_trim(row.get("title")),
        }

        display_name = (
            full_name or " ".join(part for part in [first_name, last_name] if part) if ui_type == "Individual" else full_name or last_name
        )
        queries[party_key] = _build_entity_example(item)
        row_meta.append(BatchRowMeta(key=party_key, display_name=display_name or f"(Row {index + 1})", ui_type=ui_type))

    if validation_errors:
        raise BatchUploadValidationError(validation_errors)
    if not queries:
        raise BatchUploadValidationError(["All rows are invalid (missing required names)."])

    return ParsedBatchUpload(queries=queries, row_meta=row_meta)
