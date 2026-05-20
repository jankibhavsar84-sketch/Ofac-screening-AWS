from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MD = REPO_ROOT / "docs" / "technical-design-document.md"
DEFAULT_DOCX = REPO_ROOT / "docs" / "technical-design-document.docx"


def _ensure_style(document: Document, name: str) -> str:
    try:
        document.styles[name]
        return name
    except Exception:
        return "Normal"


def _add_code_block(document: Document, lines: list[str]) -> None:
    style_name = _ensure_style(document, "CodeBlock")
    if style_name == "Normal":
        try:
            style = document.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = document.styles["Normal"]
            style.font.name = "Consolas"
            style.font.size = Pt(9)
            style_name = "CodeBlock"
        except Exception:
            style_name = "Normal"

    paragraph = document.add_paragraph(style=style_name)
    run = paragraph.add_run(lines[0] if lines else "")
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    for line in lines[1:]:
        run.add_break(WD_BREAK.LINE)
        run.add_text(line)


def _add_rich_text_paragraph(document: Document, text: str, *, style: str = "Normal") -> None:
    paragraph = document.add_paragraph(style=_ensure_style(document, style))

    # Very small markdown subset:
    # - **bold**
    # - `inline code`
    # Preserve everything else as plain text.
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    parts = token_re.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            continue
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue
        paragraph.add_run(part)


def _parse_pipe_table(lines: list[str]) -> list[list[str]]:
    # Supports simple GitHub-flavored tables:
    # | a | b |
    # |---|---|
    # | 1 | 2 |
    def split_row(row: str) -> list[str]:
        raw = row.strip().strip("|")
        return [c.strip() for c in raw.split("|")]

    rows = [split_row(line) for line in lines]
    if len(rows) >= 2 and all(re.fullmatch(r"[:\-\s]+", cell) for cell in rows[1]):
        rows.pop(1)  # separator row
    return rows


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=col_count)
    try:
        table.style = "Table Grid"
    except Exception:
        pass

    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.text = row[c_idx] if c_idx < len(row) else ""
            if r_idx == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def _add_image(document: Document, image_path: Path, alt_text: str = "") -> None:
    if not image_path.exists():
        _add_rich_text_paragraph(document, alt_text or f"Image not found: {image_path}")
        return
    document.add_picture(str(image_path), width=Inches(6.5))
    if document.paragraphs:
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def render(md_path: Path = DEFAULT_MD, docx_path: Path = DEFAULT_DOCX) -> None:
    md_text = md_path.read_text(encoding="utf-8")
    lines = md_text.splitlines()

    document = Document()

    in_code = False
    code_fence = ""
    code_lines: list[str] = []
    pending_paragraph: list[str] = []
    pending_table: list[str] = []

    def flush_paragraph() -> None:
        nonlocal pending_paragraph
        text = " ".join(s.strip() for s in pending_paragraph).strip()
        pending_paragraph = []
        if text:
            _add_rich_text_paragraph(document, text)

    def flush_table() -> None:
        nonlocal pending_table
        if not pending_table:
            return
        rows = _parse_pipe_table(pending_table)
        pending_table = []
        if rows:
            _add_table(document, rows)

    for raw in lines:
        line = raw.rstrip("\n")

        # Code fences
        if line.strip().startswith("```"):
            if not in_code:
                flush_paragraph()
                flush_table()
                in_code = True
                code_fence = line.strip()
                code_lines = []
                continue
            if in_code and (line.strip() == code_fence or line.strip() == "```"):
                _add_code_block(document, code_lines)
                in_code = False
                code_fence = ""
                code_lines = []
                continue

        if in_code:
            code_lines.append(line)
            continue

        # Tables (simple pipe tables)
        is_table_line = "|" in line and line.strip().startswith("|")
        if is_table_line:
            flush_paragraph()
            pending_table.append(line)
            continue
        flush_table()

        # Images
        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line.strip())
        if image_match:
            flush_paragraph()
            alt_text = image_match.group(1).strip()
            image_ref = image_match.group(2).strip()
            if not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", image_ref):
                image_path = Path(image_ref)
                if not image_path.is_absolute():
                    image_path = md_path.parent / image_path
                _add_image(document, image_path.resolve(), alt_text)
            else:
                _add_rich_text_paragraph(document, alt_text or image_ref)
            continue

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line.strip())
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            _add_rich_text_paragraph(document, title, style=style)
            continue

        # Bullets (supports nesting by 2-space indents)
        bullet_match = re.match(r"^(\s*)-\s+(.*)$", line)
        if bullet_match:
            flush_paragraph()
            indent = len(bullet_match.group(1).replace("\t", "  "))
            level = max(0, min(2, indent // 2))
            bullet_style = ["List Bullet", "List Bullet 2", "List Bullet 3"][level]
            _add_rich_text_paragraph(document, bullet_match.group(2).strip(), style=bullet_style)
            continue

        # Blank line = paragraph boundary
        if not line.strip():
            flush_paragraph()
            continue

        pending_paragraph.append(line.strip())

    flush_paragraph()
    flush_table()
    if in_code:
        _add_code_block(document, code_lines)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


if __name__ == "__main__":
    render()
